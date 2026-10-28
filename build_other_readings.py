from pathlib import Path
from docx import Document
import json,re,base64
sources=[('1_ Teacher','other-1-teacher',24),('2_ Teacher','other-2-teacher',15),('3_ Teacher','other-3-teacher',10),('3_Student','other-3-student',13)]
def cells(row):
 seen=set();out=[]
 for c in row.cells:
  if c._tc in seen:continue
  seen.add(c._tc);out.append(c)
 return out
def vals(row):return [c.text.strip() for c in cells(row) if c.text.strip() and not re.fullmatch(r'\d+\)?',c.text.strip())]
def islabels(s):return bool(re.search(r'A\).*B\).*C\).*D\)',s,re.S))
def keygrid(t):
 mapping={};headers={};counter=0
 for row in t.rows:
  values=[c.text.strip() for c in row.cells]
  if any(re.fullmatch(r'Text\s*\d+',v) for v in values):
   headers={i:int(re.search(r'\d+',v)[0]) for i,v in enumerate(values) if re.fullmatch(r'Text\s*\d+',v)};counter=0;continue
  if not any(re.fullmatch(r'[A-D]',v) for v in values):continue
  counter+=1
  for col,tid in headers.items():
   if col<len(values) and re.fullmatch('[A-D]',values[col]):mapping.setdefault(tid,[]).append(ord(values[col])-65)
 return mapping
collections=[];meta=[]
for suffix,cid,expected in sources:
 path=Path('upload')/('ENGL010_Supplemental Reading Tests '+suffix+'.docx');d=Document(path)
 if suffix.startswith('2_'):
  keys={i:[ord(x)-65 for x in re.findall(r'\d+\.\s*([A-D])',' '.join(c.text for c in r.cells[1:]))] for i,r in enumerate(d.tables[-1].rows[1:],1)}
 else:keys=keygrid(d.tables[-1])
 tests=[];current=None
 for ti,t in enumerate(d.tables[:-1]):
  if ti==0 and suffix.startswith('1_'):continue
  pending=''
  passageCells=[]
  # Each passage occupies the initial row(s), before the first MC stem/options pair.
  optionRows=[]
  for ri,row in enumerate(t.rows):
   v=vals(row);l=islabels('\n'.join(c.text for c in cells(row)))
   missing=bool(ri>0 and not l and len(v)==1 and len([x for x in v[0].splitlines() if x.strip()])==4 and len(vals(t.rows[ri-1]))==1 and len(vals(t.rows[ri-1])[0])<350 and re.search(r'[?]|_{3}|PURPOSE',vals(t.rows[ri-1])[0]))
   if l or missing:optionRows.append(ri)
  first=min(optionRows)-1 if optionRows else len(t.rows)
  if any(len(v)>600 for row in t.rows[:max(1,first)] for v in vals(row)):
   if current:tests.append(current)
   passage=[];seen=set()
   for row in t.rows[:max(1,first)]:
    for c in cells(row):
     txt=c.text.strip()
     if txt in seen or len(txt)<18 or re.fullmatch(r'[\d\s]+',txt):continue
     seen.add(txt);passage.append(txt)
   raw='\n\n'.join(passage)
   paras=[x.strip() for x in re.split(r'\n\s*\n',raw) if x.strip()]
   if paras and re.fullmatch(r'_+\s*\(See (?:question 1\.|the appropriate question\.)\)',paras[0],re.I):paras.pop(0)
   current={'id':len(tests)+1,'collection':cid,'title':'Reading '+str(len(tests)+1),'paragraphs':paras,'questions':[],'vocabulary':[],'sourceFile':path.name,'sourceTable':ti,'sourceRows':first,'sourceText':raw}
   meta.append({'collection':cid,'id':current['id'],'table':ti,'source':suffix,'text':raw})
  if not current:continue
  for ri,row in enumerate(t.rows):
   values=vals(row)
   # Read choice rows by their label column; recover a missing label column if four text paragraphs follow a question.
   options=[]
   labels=islabels('\n'.join(c.text for c in cells(row)))
   possible=bool(ri>0 and not labels and len(values)==1 and len([x for x in values[0].splitlines() if x.strip()])==4 and len(vals(t.rows[ri-1]))==1 and len(vals(t.rows[ri-1])[0])<350 and re.search(r'[?]|_{3}|PURPOSE',vals(t.rows[ri-1])[0]))
   if not labels and not possible:continue
   for c in cells(row):
    if c._tc.xpath('.//a:blip'):
     b=c._tc.xpath('.//a:blip')[0];part=d.part.related_parts[b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')]
     options.append({'image':'data:'+part.content_type+';base64,'+base64.b64encode(part.blob).decode()})
   if not options:
    payload=[v for v in values if not islabels(v)]
    if not payload:raise ValueError((cid,ti,ri,'NO OPTIONS'))
    options=[{'text':v.strip()} for v in payload[-1].splitlines() if v.strip()]
   if len(options)!=4:raise ValueError((cid,ti,ri,len(options),options))
   previous=vals(t.rows[ri-1]);stem=previous[-1] if previous else ''
   if not stem:raise ValueError((cid,ti,ri,'NO STEM'))
   
   stemcell=max(cells(t.rows[ri-1]),key=lambda c:len(c.text))
   target=''.join(run.text for paragraph in stemcell.paragraphs for run in paragraph.runs if run.italic).strip()
   current['questions'].append({'stem':stem,'sourceStem':stem,'target':target or None,'options':options,'key':None,'questionTable':ti,'questionRow':ri-1})
 if current:tests.append(current)
 assert len(tests)==expected,(cid,len(tests))
 for t in tests:
  assert len(t['questions'])==10,(cid,t['id'],len(t['questions']))
  assert len(keys[t['id']])==10,(cid,t['id'],keys[t['id']])
  for q,k in zip(t['questions'],keys[t['id']]):q['key']=k
 collections.append({'id':cid,'title':'Reading Tests '+suffix.replace('_ ','_'),'filename':path.name,'tests':tests})
 print(cid,len(tests),'readings',sum(len(t['questions']) for t in tests),'questions')
Path('other-reading-tests.json').write_text(json.dumps(collections,ensure_ascii=False,separators=(',',':')))
Path('other-reading-source-map.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2))
