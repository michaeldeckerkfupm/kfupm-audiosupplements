from docx import Document
from pathlib import Path
import json,re,base64
D=Document('upload/ENGL010_Supplemental Reading Tests 5_Student.docx')
titles=[p.text.split(': ',1)[1] for p in D.paragraphs if re.match(r'Reading \d+:',p.text)][:18]
def unique_cells(row):
 seen=set();out=[]
 for c in row.cells:
  if c._tc in seen:continue
  seen.add(c._tc);out.append(c)
 return out
def text(row):return [c.text.strip() for c in unique_cells(row) if c.text.strip()]
def stem(s):
 s=re.sub(r'\s*\(Lines? [\d\s–—-]+\)','',s,flags=re.I)
 s=re.sub(r'^In Line \d+,\s*','',s,flags=re.I)
 s=re.sub(r' in Line \d+','',s,flags=re.I)
 return s[0].upper()+s[1:]
starts=[[1,8,12,19],[1,6,12,21],[1,12,24],[1,9,16,24],[1,9,17,25,33],[1,11,20,26],[1,10,17],[1,8,14,23,32],[1,11,19,25],[1,10,14,22,29],[1,5,13,17,24],[1,7,12,20],[1,7,14,21],[1,5,15,23,33],[1,7,17,30],[1,8,13,24,35],[1,7,15,26,33],[1,5,13,21,29]]
tests=[]
for i,t in enumerate(D.tables[:18]):
 passage=max(text(t.rows[0]),key=len)
 # Keep source paragraph breaks and gap placeholders. The initial title gap is not a paragraph.
 paras=[x.strip() for x in passage.split('\n\n') if x.strip()]
 if re.fullmatch(r'_+\s*\(See question 1\.\)',paras[0]):paras.pop(0)
 qs=[]
 for n in range(2,len(t.rows),2):
  options=[]
  cells=unique_cells(t.rows[n+1])
  for c in cells:
   if c._tc.xpath('.//a:blip'):
    b=c._tc.xpath('.//a:blip')[0]
    rel=b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    part=D.part.related_parts[rel]
    options.append({'image':'data:'+part.content_type+';base64,'+base64.b64encode(part.blob).decode()})
  if not options:
   options=[{'text':v.strip()} for v in text(t.rows[n+1])[-1].splitlines() if v.strip()]
  assert len(options)==4,(i,n,options)
  answer=text(D.tables[18+i].rows[len(qs)+1])[-1]
  key=ord(answer[0])-65
  assert key in range(4),(i,answer)
  original=text(t.rows[n])[-1]
  target=re.search(r'Line (\d+)',original)
  reference=max(j+1 for j,start in enumerate(starts[i]) if start<=int(target[1])) if target else None
  if i==7 and len(qs)==9:
   for o in options:
    line=int(re.search(r'\d+',o['text'])[0]);o['text']='Paragraph '+str(max(j+1 for j,start in enumerate(starts[i]) if start<=line))
  qs.append({'stem':stem(original),'options':options,'key':key,'reference':reference})
 assert len(qs)==len(D.tables[18+i].rows)-1
 vocab=[dict(zip(['word','location','level','pos','example','arabic'],[c.text for c in r.cells])) for r in D.tables[36+i].rows[1:]]
 tests.append({'id':i+1,'title':titles[i],'paragraphs':paras,'questions':qs,'vocabulary':vocab})
# Clarify one source item whose original options A and C were both defensible.
tests[15]['questions'][1].update(stem='According to the passage, the first recorded example of piracy ______.',options=[{'text':'took place in the Caribbean'},{'text':'took place in the South China Sea'},{'text':'took place in the Mediterranean 3,400 years ago'},{'text':'took place in the Indian Ocean'}],key=2)
Path('reading-tests.json').write_text(json.dumps(tests,ensure_ascii=False,separators=(',',':')))
print('Readings:',len(tests),'Questions:',sum(len(t['questions']) for t in tests),'Images:',sum('image' in o for t in tests for q in t['questions'] for o in q['options']))
